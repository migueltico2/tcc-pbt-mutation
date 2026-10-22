#!/usr/bin/env python3
"""Gera a planilha de revisão da classificação dos mutantes não detectados.

Entrada: analysis/resultados/sobreviventes_classificados.csv (classificar_sobreviventes.py)
Saída:   analysis/revisao/revisao_sobreviventes.xlsx

A planilha mostra a pré-classificação e a justificativa de cada mutante; o revisor preenche a coluna
"Classe revisada" (lista suspensa) e, se quiser, um comentário. A aba Resumo conta as classes finais
(a revisada quando preenchida; senão, a pré-classificação) e a concordância. Depois da revisão,
analysis/importar_revisao.py leva as classes finais de volta para analysis/resultados/.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
ENT = RAIZ / "analysis" / "resultados" / "sobreviventes_classificados.csv"
SAIDA = RAIZ / "analysis" / "revisao" / "revisao_sobreviventes.xlsx"
CLASSES = ["entrada não gerada", "oráculo fraco", "regra sem propriedade", "equivalente"]
DEFINICOES = {
    "entrada não gerada": "O oráculo detecta o defeito, mas só quando o gerador produz a entrada que o revela "
                          "(em geral um valor exato de limite). Inclui mutantes detectados em parte das execuções.",
    "oráculo fraco": "Nunca detectado, embora a regra seja verificada por alguma propriedade da suíte: a propriedade "
                     "exercita o trecho mas não observa a diferença.",
    "regra sem propriedade": "Nunca detectado, e a regra que o trecho implementa não é verificada por nenhuma "
                             "propriedade da suíte.",
    "equivalente": "Nenhuma entrada permitida pelo SPEC distingue o mutante do original na API de domínio.",
}

FONTE = "Arial"
F = Font(name=FONTE, size=10)
FB = Font(name=FONTE, size=10, bold=True)
FT = Font(name=FONTE, size=14, bold=True)
FH = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
CAB = PatternFill("solid", start_color="2A4D69")
ENTRADA = PatternFill("solid", start_color="FFF2CC")       # células para o revisor preencher
CINZA = PatternFill("solid", start_color="F2F2F2")
DIVERGE = PatternFill("solid", start_color="F8CBAD")
FINO = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
QUEBRA = Alignment(wrap_text=True, vertical="top")


def main():
    df = pd.read_csv(ENT, keep_default_na=False)
    df = df.sort_values(["ecossistema", "suite", "classe", "arquivo", "linha"]).reset_index(drop=True)
    wb = Workbook()

    # ---------------------------------------------------------------- Leia-me
    ws = wb.active
    ws.title = "Leia-me"
    ws["A1"] = "Revisão da classificação dos mutantes não detectados"
    ws["A1"].font = FT
    linhas = [
        "Cada linha da aba Mutantes é um mutante distinto, em uma combinação de ecossistema e suíte, que não foi "
        "detectado em ao menos uma das 90 execuções (3 níveis de tentativas × 30).",
        "As colunas até \"Justificativa (pré-classificação)\" vêm da análise automática e da leitura do código. "
        "Não é preciso alterá-las.",
        "O que preencher: só as células em amarelo, na aba Mutantes. Em \"Classe revisada\", escolha uma classe na "
        "lista suspensa SOMENTE se discordar da pré-classificação ou quiser confirmá-la explicitamente. "
        "Em \"Comentário do revisor\", escreva o motivo quando discordar.",
        "A coluna \"Classe final\" usa a classe revisada quando preenchida e, senão, a pré-classificação. "
        "A coluna \"Concorda?\" fica vermelha quando há divergência. A aba Resumo atualiza sozinha.",
        "A linha 2 da aba Mutantes traz um exemplo de preenchimento (em cinza) e não entra nas contagens.",
    ]
    for i, t in enumerate(linhas, start=3):
        ws.cell(row=i, column=1, value=t).font = F
        ws.cell(row=i, column=1).alignment = QUEBRA
    ws.cell(row=9, column=1, value="Definição das classes").font = FB
    for i, c in enumerate(CLASSES, start=10):
        ws.cell(row=i, column=1, value=f"{c}: {DEFINICOES[c]}").font = F
        ws.cell(row=i, column=1).alignment = QUEBRA
    ws.column_dimensions["A"].width = 120
    for r in range(3, 14):
        ws.row_dimensions[r].height = 32

    # ---------------------------------------------------------------- Mutantes
    wm = wb.create_sheet("Mutantes")
    colunas = [
        ("Nº", 6), ("Ecossistema", 11), ("Suíte", 8), ("Arquivo", 20), ("Linha", 7), ("Código", 46),
        ("Mutação", 40), ("Regra", 9), ("Detectado em (de 90)", 11), ("Taxa t=10", 8), ("Taxa t=100", 8),
        ("Taxa t=1000", 8), ("Suíte forte detecta?", 10), ("Pré-classificação", 18),
        ("Justificativa (pré-classificação)", 48), ("Classe revisada", 18), ("Comentário do revisor", 36),
        ("Classe final", 18), ("Concorda?", 10), ("ID do mutante (não editar)", 30),
    ]
    for j, (nome, larg) in enumerate(colunas, start=1):
        c = wm.cell(row=1, column=j, value=nome)
        c.font, c.fill, c.alignment, c.border = FH, CAB, Alignment(wrap_text=True, vertical="center"), BORDA
        wm.column_dimensions[get_column_letter(j)].width = larg
    wm.row_dimensions[1].height = 34
    wm.freeze_panes = "D2"

    # linha 2: exemplo (não conta)
    ex = ["ex.", "java", "fraca", "CalculadoraPreco.java", 61, "if (subtotal < cupom.getMinimoCentavos()) {",
          "changed conditional boundary", "R12", 0, 0, 0, 0, "sim", "entrada não gerada",
          "o defeito só aparece com subtotal igual ao mínimo do cupom", "entrada não gerada",
          "concordo: P9 pegaria se o gerador produzisse subtotal = mínimo"]
    for j, v in enumerate(ex, start=1):
        c = wm.cell(row=2, column=j, value=v)
        c.font, c.fill, c.alignment, c.border = Font(name=FONTE, size=10, italic=True, color="7F7F7F"), CINZA, QUEBRA, BORDA
    wm.cell(row=2, column=18, value="=IF(P2=\"\",N2,P2)").font = Font(name=FONTE, size=10, italic=True, color="7F7F7F")
    wm.cell(row=2, column=19, value="=IF(P2=\"\",\"\",IF(P2=N2,\"sim\",\"não\"))").font = Font(name=FONTE, size=10, italic=True, color="7F7F7F")

    primeira = 3
    for k, r in enumerate(df.itertuples(), start=0):
        i = primeira + k
        valores = [k + 1, r.ecossistema, r.suite, r.arquivo, int(r.linha), r.codigo, r.mutacao, r.regra,
                   int(r.detectado_em), float(r.taxa_deteccao_t10), float(r.taxa_deteccao_t100),
                   float(r.taxa_deteccao_t1000), "sim" if int(r.suite_forte_detecta) else "não", r.classe,
                   r.justificativa, None, None]
        for j, v in enumerate(valores, start=1):
            c = wm.cell(row=i, column=j, value=v)
            c.font, c.alignment, c.border = F, QUEBRA, BORDA
            if j in (10, 11, 12):
                c.number_format = "0%"
            if j in (16, 17):
                c.fill = ENTRADA
        wm.cell(row=i, column=18, value=f'=IF(P{i}="",N{i},P{i})').font = F
        wm.cell(row=i, column=19, value=f'=IF(P{i}="","",IF(P{i}=N{i},"sim","não"))').font = F
        wm.cell(row=i, column=20, value=r.mutante_id).font = Font(name=FONTE, size=8, color="7F7F7F")
        for j in (18, 19, 20):
            wm.cell(row=i, column=j).border = BORDA
            wm.cell(row=i, column=j).alignment = QUEBRA
    ultima = primeira + len(df) - 1

    dv = DataValidation(type="list", formula1='"' + ",".join(CLASSES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Classe inválida",
                        error="Escolha uma das quatro classes da lista.")
    wm.add_data_validation(dv)
    dv.add(f"P{primeira}:P{ultima}")
    wm.conditional_formatting.add(f"S{primeira}:S{ultima}",
                                  FormulaRule(formula=[f'$S{primeira}="não"'], fill=DIVERGE))
    wm.auto_filter.ref = f"A1:T{ultima}"
    wm["P1"].comment = Comment("Preencha só se discordar (ou para confirmar). Use a lista suspensa.", "Revisão")

    # ---------------------------------------------------------------- Resumo
    wr = wb.create_sheet("Resumo")
    wr["A1"] = "Classes finais por ecossistema e suíte (atualiza com a revisão)"
    wr["A1"].font = FT
    cab = ["Ecossistema", "Suíte"] + CLASSES + ["Total", "Revisados", "Divergências"]
    for j, nome in enumerate(cab, start=1):
        c = wr.cell(row=3, column=j, value=nome)
        c.font, c.fill, c.border = FH, CAB, BORDA
        c.alignment = Alignment(wrap_text=True, vertical="center")
        wr.column_dimensions[get_column_letter(j)].width = 16 if j > 2 else 13
    rng = lambda col: f"Mutantes!${col}${primeira}:${col}${ultima}"
    linha = 4
    for eco in ["java", "node"]:
        for suite in ["fraca", "forte"]:
            wr.cell(row=linha, column=1, value=eco).font = F
            wr.cell(row=linha, column=2, value=suite).font = F
            for j, classe in enumerate(CLASSES, start=3):
                col = get_column_letter(j)
                wr.cell(row=linha, column=j,
                        value=f'=COUNTIFS({rng("B")},$A{linha},{rng("C")},$B{linha},{rng("R")},{col}$3)').font = F
            wr.cell(row=linha, column=7, value=f"=SUM(C{linha}:F{linha})").font = FB
            wr.cell(row=linha, column=8,
                    value=f'=COUNTIFS({rng("B")},$A{linha},{rng("C")},$B{linha},{rng("P")},"<>")').font = F
            wr.cell(row=linha, column=9,
                    value=f'=COUNTIFS({rng("B")},$A{linha},{rng("C")},$B{linha},{rng("S")},"não")').font = F
            for j in range(1, 10):
                wr.cell(row=linha, column=j).border = BORDA
            linha += 1
    wr.cell(row=linha + 1, column=1, value="Concordância entre a pré-classificação e a revisão (só mutantes revisados):").font = FB
    wr.cell(row=linha + 2, column=1, value="Revisados").font = F
    wr.cell(row=linha + 2, column=2, value=f"=SUM(H4:H{linha - 1})").font = F
    wr.cell(row=linha + 3, column=1, value="Concordância").font = F
    wr.cell(row=linha + 3, column=2, value=f"=IFERROR(1-SUM(I4:I{linha - 1})/B{linha + 2},\"\")").font = F
    wr.cell(row=linha + 3, column=2).number_format = "0.0%"

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SAIDA)
    print(f"{len(df)} mutantes -> {SAIDA}")


if __name__ == "__main__":
    main()
